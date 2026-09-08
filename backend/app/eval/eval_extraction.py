"""
Extraction accuracy evaluation against a human-verified ground truth file.

This is the check the per-run "verification report" cannot do. That report
reopens the saved Excel and diffs it against the extracted JSON -- both sides
come from the same extraction, so a misread GSTIN passes. Here the reference is
a ground truth file transcribed by a human from the source documents, so a
misread GSTIN shows up as `wrong`.

Ground truth can come from either source:

    python -m eval.eval_extraction \\
        --ground-truth eval/ground_truth/mb_control_systems.yaml \\
        --documents "path/to/vendor/docs"

    python -m eval.eval_extraction \\
        --ground-truth-excel "path/to/vendor/docs/GROUND_TRUTH.xlsx" \\
        --mapping vendor_creation_v1 \\
        --documents "path/to/vendor/docs"

The Excel form is a pre-filled copy of the same template the pipeline writes
into (e.g. VENDOR CREATION REQUEST FORM.xlsx), hand-filled by a human with the
correct value in every cell the mapping writes. Because it uses the same cell
map as the pipeline's own output, comparison is a direct cell-for-cell read via
`ExcelMapper.cells_for()` -- no separate transcription format to keep in sync.
Blank cells are treated as `absent: true` (nothing on this vendor's ground-truth
form), not as "not yet filled in" -- an incompletely-filled ground-truth sheet
will silently under-report `missed`/`wrong` on the fields left blank, so treat a
"0 wrong" run with suspicion until every cell has actually been filled in.

Outcomes per field:

    correct           extracted == ground truth
    wrong             ground truth has a value, extractor produced a DIFFERENT one
    missed            ground truth has a value, extractor produced nothing
    hallucinated      ground truth says absent, extractor produced a value
    correctly_absent  ground truth says absent, extractor produced nothing

`wrong` and `hallucinated` are the ones that matter most: they put a
confident, well-formed, incorrect value on a form a human is likely to sign.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
from pathlib import Path
from typing import Any, Optional

os.environ.setdefault("PADDLE_PDX_ENABLE_MKLDNN_BYDEFAULT", "0")

import yaml

GROUND_TRUTH_DIR = Path(__file__).resolve().parent / "ground_truth"

# Identifiers must match exactly -- a single wrong character makes them wrong,
# and "close enough" is precisely the failure mode being measured. Free text is
# compared after normalization, since capitalisation and punctuation differences
# between "M B CONTROL" and "M.B. CONTROL" are formatting, not extraction errors.
EXACT_MATCH_FIELDS = {
    "gst_number",
    "pan",
    "tan",
    "udyam_number",
    "esic_number",
    "ifsc",
    "account_number",
    "pin_code",
    "telephone",
    "email",
}

_PUNCT = re.compile(r"[.,\-/&()]+")
_SPACE = re.compile(r"\s+")

# The address is filled as one unit -- address_1..4 + city + state + pin_code.
# A combined "Address of Principal Place of Business" line that the resolver
# fails to segment shows up as several independent `missed`/`wrong` fields;
# grouping them into one block verdict (below) makes an address regression
# legible as a single number instead of being averaged away across 24 fields.
ADDRESS_FIELDS = (
    "address_1", "address_2", "address_3", "address_4",
    "city", "state", "pin_code",
)


class GroundTruthError(ValueError):
    pass


# ---------------------------------------------------------------------------
# GROUND TRUTH
# ---------------------------------------------------------------------------

def load_ground_truth(path: Path, allow_unverified: bool = False) -> dict:
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    if "fields" not in data:
        raise GroundTruthError(f"{path.name}: no 'fields' block")

    if not data.get("verified"):
        message = (
            f"{path.name} is marked verified: false -- its values have not been "
            f"confirmed against the source documents by a human, so any score "
            f"computed from it is meaningless."
        )
        if not allow_unverified:
            raise GroundTruthError(message + " Pass --allow-unverified to score anyway.")
        print(f"\n  !! WARNING: {message}\n")

    for key, spec in data["fields"].items():
        if not isinstance(spec, dict):
            raise GroundTruthError(f"field {key!r} must be a mapping")
        if spec.get("absent") and "value" in spec:
            raise GroundTruthError(f"field {key!r} is marked absent but also has a value")
        if not spec.get("absent") and "value" not in spec:
            raise GroundTruthError(f"field {key!r} has neither a value nor absent: true")

    return data


def load_ground_truth_excel(path: Path, mapping_name: str, vendor_id: Optional[str] = None) -> dict:
    """Read a hand-filled copy of the pipeline's own Excel template as ground
    truth. Same cell map as ExcelMapper.fill() / verify_excel() use, so a value
    in cell C7 here is compared against exactly the field the pipeline writes
    into C7 -- no separate field-name mapping to maintain.

    Unlike the YAML path there is no `verified: true` flag: filling in every
    cell by hand *is* the verification step. A blank cell means "nothing to
    extract here" (`absent: true`), not "not filled in yet" -- see the module
    docstring's warning about incompletely-filled sheets.
    """
    import openpyxl

    from ..services.extraction_pipeline.excel.excel_mapper import ExcelMapper

    mapper = ExcelMapper.load(mapping_name)
    workbook = openpyxl.load_workbook(str(path), data_only=True)

    fields: dict[str, dict] = {}
    for field, cell_ref, sheet in mapper.cells_for():
        target = sheet or workbook.sheetnames[0]
        if target not in workbook.sheetnames:
            raise GroundTruthError(
                f"{path.name}: sheet {target!r} not found. Available: {workbook.sheetnames}"
            )
        value = workbook[target][cell_ref].value
        if value is None or str(value).strip() == "":
            fields[field] = {"absent": True}
        else:
            fields[field] = {"value": value}

    return {"vendor_id": vendor_id or path.stem, "verified": True, "fields": fields}


# ---------------------------------------------------------------------------
# COMPARISON
# ---------------------------------------------------------------------------

def _normalize(value: Any, exact: bool) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    # openpyxl/JSON round-trips can turn "700019" into 700019.0
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    if exact:
        return text.upper().replace(" ", "")
    text = _PUNCT.sub(" ", text)
    text = _SPACE.sub(" ", text)
    return text.strip().lower()


def compare_field(key: str, gt_spec: dict, extracted: Any) -> dict:
    exact = key in EXACT_MATCH_FIELDS
    has_extracted = extracted is not None and str(extracted).strip() != ""

    if gt_spec.get("absent"):
        outcome = "hallucinated" if has_extracted else "correctly_absent"
        return {
            "field": key,
            "outcome": outcome,
            "expected": None,
            "extracted": extracted if has_extracted else None,
        }

    expected = gt_spec["value"]
    if not has_extracted:
        outcome = "missed"
    elif _normalize(expected, exact) == _normalize(extracted, exact):
        outcome = "correct"
    else:
        outcome = "wrong"

    return {
        "field": key,
        "outcome": outcome,
        "expected": expected,
        "extracted": extracted if has_extracted else None,
    }


def evaluate(ground_truth: dict, extracted: dict) -> list[dict]:
    return [
        compare_field(key, spec, extracted.get(key))
        for key, spec in ground_truth["fields"].items()
    ]


def _address_block(results: list[dict]) -> Optional[dict]:
    """Roll the address fields up into one verdict. None when the ground truth
    carries no address fields at all."""
    rows = [r for r in results if r["field"] in ADDRESS_FIELDS]
    if not rows:
        return None
    counts = {"correct": 0, "wrong": 0, "missed": 0,
              "hallucinated": 0, "correctly_absent": 0}
    for r in rows:
        counts[r["outcome"]] += 1
    bad = counts["wrong"] + counts["missed"] + counts["hallucinated"]
    scored = len(rows) - counts["correctly_absent"]
    return {
        **counts,
        "fields": len(rows),
        # every field the GT expects a value for was got right
        "exact": bad == 0,
        "field_accuracy": round(counts["correct"] / scored * 100, 1) if scored else None,
    }


def summarize(results: list[dict]) -> dict:
    counts = {
        "correct": 0, "wrong": 0, "missed": 0,
        "hallucinated": 0, "correctly_absent": 0,
    }
    for r in results:
        counts[r["outcome"]] += 1

    # Fields where ground truth says a value exists.
    present = counts["correct"] + counts["wrong"] + counts["missed"]
    # Fields where the extractor produced something.
    produced = counts["correct"] + counts["wrong"] + counts["hallucinated"]

    return {
        **counts,
        "total": len(results),
        "recall": round(counts["correct"] / present * 100, 1) if present else None,
        "precision": round(counts["correct"] / produced * 100, 1) if produced else None,
        "accuracy": round(
            (counts["correct"] + counts["correctly_absent"]) / len(results) * 100, 1
        ) if results else None,
        "address_block": _address_block(results),
    }


# ---------------------------------------------------------------------------
# EXTRACTORS
# ---------------------------------------------------------------------------

def extract_v2(document_dir: Path, models: str = "small", cache: Optional[Path] = None) -> dict:
    from app.services.extraction_pipeline.ingest.document_loader import load_documents
    from app.services.extraction_pipeline.models import DocumentSet
    from app.services.extraction_pipeline.ingest.ocr_engine import OCREngine
    from app.services.extraction_pipeline.pipeline import collect_inputs, extract_from_document_set

    if cache and cache.exists():
        print(f"  using cached document set: {cache}")
        doc_set = DocumentSet.load_json(str(cache))
    else:
        files = collect_inputs([document_dir])
        if not files:
            raise SystemExit(f"No supported documents found in {document_dir}")
        print(f"  loading {len(files)} document(s) (OCR, this takes a couple of minutes)...")
        engine = OCREngine(
            det_model=f"PP-OCRv6_{models}_det",
            rec_model=f"PP-OCRv6_{models}_rec",
        )
        doc_set = load_documents(files, engine=engine)
        if cache:
            cache.parent.mkdir(parents=True, exist_ok=True)
            doc_set.save_json(str(cache))
            print(f"  cached document set -> {cache}")

    return extract_from_document_set(doc_set).canonical()


# ---------------------------------------------------------------------------
# REPORTING
# ---------------------------------------------------------------------------

_ICON = {
    "correct": "OK  ",
    "wrong": "WRONG",
    "missed": "MISS",
    "hallucinated": "HALLU",
    "correctly_absent": "-   ",
}


def print_report(results: list[dict], summary: dict, ground_truth: dict) -> None:
    print("\n" + "=" * 92)
    print(f"EXTRACTION ACCURACY vs GROUND TRUTH ({ground_truth.get('vendor_id')})")
    print("=" * 92)
    print(f"  {'':<5} {'field':<20} {'expected':<30} extracted")
    print(f"  {'-'*5} {'-'*20} {'-'*30} {'-'*28}")

    order = ["wrong", "hallucinated", "missed", "correct", "correctly_absent"]
    for outcome in order:
        for r in (x for x in results if x["outcome"] == outcome):
            exp = str(r["expected"])[:29] if r["expected"] is not None else "(absent)"
            got = str(r["extracted"])[:28] if r["extracted"] is not None else "(none)"
            print(f"  {_ICON[outcome]:<5} {r['field']:<20} {exp:<30} {got}")

    if ground_truth.get("conflicts"):
        print("\n" + "-" * 92)
        print("DOCUMENT CONFLICTS (ground truth took a position; verify it matches your process)")
        print("-" * 92)
        for c in ground_truth["conflicts"]:
            note = " ".join(str(c.get("note", "")).split())
            print(f"  {c.get('field')}: {note}")

    print("\n" + "=" * 92)
    print("SUMMARY")
    print("=" * 92)
    print(f"  correct          : {summary['correct']}")
    print(f"  wrong            : {summary['wrong']}      <- confident but incorrect")
    print(f"  missed           : {summary['missed']}")
    print(f"  hallucinated     : {summary['hallucinated']}      <- invented a value that isn't on any document")
    print(f"  correctly absent : {summary['correctly_absent']}")
    print(f"  {'-'*40}")
    print(f"  fields           : {summary['total']}")
    print(f"  recall           : {summary['recall']}%   (of fields that DO have a value, how many were got right)")
    print(f"  precision        : {summary['precision']}%   (of values produced, how many were right)")
    print(f"  accuracy         : {summary['accuracy']}%   (all fields, including correctly-left-blank)")

    ab = summary.get("address_block")
    if ab:
        verdict = "PASS" if ab["exact"] else "FAIL"
        print(f"  {'-'*40}")
        print(f"  address block    : {verdict}   "
              f"({ab['correct']} correct / {ab['wrong']} wrong / {ab['missed']} missed "
              f"/ {ab['hallucinated']} hallucinated across {ab['fields']} address fields)")
        print(f"                     field accuracy {ab['field_accuracy']}%   "
              f"-- address_1..4 + city + state + pin_code scored as one unit")
    print("=" * 92)


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main() -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("--ground-truth", help="Path to a ground truth YAML file")
    parser.add_argument("--ground-truth-excel",
                        help="Path to a hand-filled copy of the pipeline's Excel template, "
                             "used as ground truth instead of a YAML file")
    parser.add_argument("--mapping", default="vendor_creation_v1",
                        help="Excel mapping config name (required with --ground-truth-excel)")
    parser.add_argument("--documents", required=True, help="Directory holding that vendor's documents")
    parser.add_argument("--models", choices=["small", "medium", "tiny"], default="small")
    parser.add_argument("--cache", help="Cache/reuse the OCR'd document set at this path")
    parser.add_argument("--allow-unverified", action="store_true",
                        help="Score against a YAML ground truth file still marked verified: false")
    parser.add_argument("--json-out", help="Write the per-field results to this JSON file")
    args = parser.parse_args()

    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, OSError):
        pass

    if bool(args.ground_truth) == bool(args.ground_truth_excel):
        parser.error("pass exactly one of --ground-truth or --ground-truth-excel")

    if args.ground_truth_excel:
        ground_truth = load_ground_truth_excel(Path(args.ground_truth_excel), args.mapping)
    else:
        ground_truth = load_ground_truth(Path(args.ground_truth), args.allow_unverified)
    document_dir = Path(args.documents)

    print(f"Running extraction over {document_dir}...")
    extracted = extract_v2(document_dir, args.models, Path(args.cache) if args.cache else None)

    results = evaluate(ground_truth, extracted)
    summary = summarize(results)
    print_report(results, summary, ground_truth)

    if args.json_out:
        out = Path(args.json_out)
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(
            json.dumps(
                {"vendor_id": ground_truth.get("vendor_id"),
                 "summary": summary, "results": results},
                indent=2, ensure_ascii=False),
            encoding="utf-8",
        )
        print(f"\n  -> {out}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
